import re
from pathlib import Path

import openpyxl


SOURCE = Path(r"C:\Users\TU\Desktop\关卡表_1-22.xlsx")
TARGET_DIR = Path(r"C:\UGit\TuyouGameJam\Assets\GameData\Configuration\Levels")
EXPORT_LEVELS = range(1, 21)

X_CENTERS = {
    "x1": 1.0 / 6.0,
    "x2": 0.5,
    "x3": 5.0 / 6.0,
}
ENEMY_IDS = {"小鸡": 0, "鸡": 1, "大公鸡": 2, "坤坤": 3}
BASKETBALL_CONFIG_ID = 3
GOOSE_CAGE_CONFIG_ID = 4


def as_number(value):
    if value is None:
        return None
    text = str(value).strip()
    try:
        number = float(text)
    except ValueError:
        return None
    return int(number) if number.is_integer() else number


def parse_cell(value):
    if value is None or not str(value).strip():
        return []

    text = re.sub(r"\s+", "", str(value).strip())

    # 鸡和篮球只读取单元格最前面的一个单位，数量后缀和组合后半段忽略。
    for name in ("大公鸡", "小鸡", "坤坤", "鸡"):
        if text.startswith(name):
            return [{
                "kind": "enemy",
                "configId": ENEMY_IDS[name],
            }]

    if text.startswith("篮球"):
        return [{
            "kind": "prop",
            "configId": BASKETBALL_CONFIG_ID,
        }]

    # 鹅笼后缀数字按工程统一配置读取，具体数字忽略。
    if text.startswith("鹅笼"):
        return [{
            "kind": "prop",
            "configId": GOOSE_CAGE_CONFIG_ID,
        }]

    match = re.fullmatch(r"门(-?\d+)", text)
    if match:
        return [{
            "kind": "additive_gate",
            "initialValue": int(match.group(1)),
        }]

    match = re.fullmatch(r"武器架/(弓箭|法杖)/(\d+)", text)
    if match:
        return [{
            "kind": "prop",
            "configId": 1 if match.group(1) == "弓箭" else 2,
        }]

    match = re.fullmatch(r"(火门|冰门|雷门)/?(\d+)", text)
    if match:
        return [{
            "kind": "element_gate",
            "elementType": {"火门": 1, "冰门": 2, "雷门": 3}[match.group(1)],
            "maxHp": int(match.group(2)),
        }]

    raise ValueError(f"无法识别单元格内容：{text}")


def parse_level(sheet, level_number):
    enemies = []
    gates = []
    props = []
    order = 0

    for row_index in range(3, sheet.max_row + 1):
        level_id = as_number(sheet.cell(row_index, 1).value)
        time_value = as_number(sheet.cell(row_index, 2).value)
        if level_id != level_number or time_value is None:
            continue

        if int(time_value) != time_value or time_value < 0:
            raise ValueError(
                f"第{level_number}关第{row_index}行的时间不是非负整秒：{time_value}"
            )

        for column_index, column in enumerate(("x1", "x2", "x3"), start=3):
            parsed_items = parse_cell(sheet.cell(row_index, column_index).value)
            for parsed in parsed_items:
                base = {
                    "spawnTime": int(time_value),
                    "spawnPosition": X_CENTERS[column],
                    "row": row_index,
                    "column": column,
                    "order": order,
                }
                order += 1

                if parsed["kind"] == "enemy":
                    enemies.append({
                        **base,
                        "configId": parsed["configId"],
                    })
                elif parsed["kind"] == "prop":
                    props.append({
                        **base,
                        "configId": parsed["configId"],
                    })
                elif parsed["kind"] == "additive_gate":
                    gates.append({
                        **base,
                        "gateType": 0,
                        "initialValue": parsed["initialValue"],
                        "elementType": 0,
                        "maxHp": 0,
                    })
                elif parsed["kind"] == "element_gate":
                    gates.append({
                        **base,
                        "gateType": 1,
                        "initialValue": 0,
                        "elementType": parsed["elementType"],
                        "maxHp": parsed["maxHp"],
                    })

    sort_key = lambda item: (
        item["spawnTime"],
        item["row"],
        item["column"],
        item["order"],
    )
    enemies.sort(key=sort_key)
    gates.sort(key=sort_key)
    props.sort(key=sort_key)
    return enemies, gates, props


def format_float(value):
    return f"{value:.6f}".rstrip("0").rstrip(".")


def write_asset(level_number, enemies, gates, props):
    level_id = level_number - 1
    target = TARGET_DIR / f"Level_{level_number:03d}.asset"
    unlocked_level_ids = [level_id + 1] if level_number < max(EXPORT_LEVELS) else []
    contains_ikun = any(entry["configId"] == ENEMY_IDS["坤坤"] for entry in enemies)
    has_element_gate = any(entry["gateType"] == 1 for entry in gates)

    lines = [
        "%YAML 1.1",
        "%TAG !u! tag:unity3d.com,2011:",
        "--- !u!114 &11400000",
        "MonoBehaviour:",
        "  m_ObjectHideFlags: 0",
        "  m_CorrespondingSourceObject: {fileID: 0}",
        "  m_PrefabInstance: {fileID: 0}",
        "  m_PrefabAsset: {fileID: 0}",
        "  m_GameObject: {fileID: 0}",
        "  m_Enabled: 1",
        "  m_EditorHideFlags: 0",
        "  m_Script: {fileID: 11500000, guid: f1ef7e0d1b7621344aee412dfe246daf, type: 3}",
        f"  m_Name: Level_{level_number:03d}",
        "  m_EditorClassIdentifier: ",
        f"  levelId: {level_id}",
        f"  displayName: {level_number}",
        "  unlockedLevelIds:",
    ]
    lines.extend(f"  - {value}" for value in unlocked_level_ids)
    lines.extend([
        "  spawnY: 8",
        "  enemyApproachY: 0",
        "  despawnY: -9",
        "  bulletDespawnY: 3",
        f"  elementDurationSecondsPerDamage: {1 if has_element_gate else 0}",
        f"  ikunBasketballConfigId: {BASKETBALL_CONFIG_ID if contains_ikun else 0}",
        "  enemySpawns:",
    ])
    for entry in enemies:
        lines.extend([
            f"  - spawnTime: {format_float(entry['spawnTime'])}",
            f"    spawnPosition: {format_float(entry['spawnPosition'])}",
            f"    configId: {entry['configId']}",
        ])

    lines.append("  gateSpawns:")
    for entry in gates:
        lines.extend([
            f"  - spawnTime: {format_float(entry['spawnTime'])}",
            f"    spawnPosition: {format_float(entry['spawnPosition'])}",
            f"    gateType: {entry['gateType']}",
            f"    initialValue: {entry['initialValue']}",
            f"    elementType: {entry['elementType']}",
            f"    maxHp: {entry['maxHp']}",
        ])

    lines.append("  propSpawns:")
    for entry in props:
        lines.extend([
            f"  - spawnTime: {format_float(entry['spawnTime'])}",
            f"    spawnPosition: {format_float(entry['spawnPosition'])}",
            f"    configId: {entry['configId']}",
        ])

    target.write_text("\n".join(lines) + "\n", encoding="utf-8")


def validate_entries(level_number, entries, name):
    previous_time = -1
    for entry in entries:
        if entry["spawnTime"] < previous_time:
            raise ValueError(f"第{level_number}关{name}时间未排序")
        if entry["spawnTime"] != int(entry["spawnTime"]):
            raise ValueError(f"第{level_number}关{name}存在非整秒时间")
        if entry["spawnPosition"] not in X_CENTERS.values():
            raise ValueError(f"第{level_number}关{name}存在非固定横向坐标")
        previous_time = entry["spawnTime"]


def validate_level(level_number, enemies, gates, props):
    if not enemies:
        raise ValueError(f"第{level_number}关没有鸡配置")
    validate_entries(level_number, enemies, "enemySpawns")
    validate_entries(level_number, gates, "gateSpawns")
    validate_entries(level_number, props, "propSpawns")


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(f"找不到关卡表：{SOURCE}")

    workbook = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
    sheet = workbook["Sheet1"]
    for level_number in EXPORT_LEVELS:
        enemies, gates, props = parse_level(sheet, level_number)
        validate_level(level_number, enemies, gates, props)
        write_asset(level_number, enemies, gates, props)
        print(
            f"level={level_number:02d} "
            f"enemySpawns={len(enemies)} "
            f"gateSpawns={len(gates)} "
            f"propSpawns={len(props)} "
            f"basketballs={sum(entry['configId'] == BASKETBALL_CONFIG_ID for entry in props)} "
            f"gooseCages={sum(entry['configId'] == GOOSE_CAGE_CONFIG_ID for entry in props)} "
            f"ikun={sum(entry['configId'] == ENEMY_IDS['坤坤'] for entry in enemies)}"
        )


if __name__ == "__main__":
    main()
